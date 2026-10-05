# -*- coding: utf-8 -*-
"""
Genera las plantillas de cliente de la fase 1 de la 5.11 (salida a Excel de
verdad, D3; docs/dev/plan_5_11_procesos_que_corren_solos.md). Los archivos
generados se versionan: este script solo hace falta para rehacerlos.

    python scripts/fixtures/excel/generar_plantillas.py

Necesita openpyxl. Cada plantilla lleva lo que un cliente tiene en su libro y
AmoxSQL no debe romper al rellenarlo: formulas, una tabla con columna
calculada y fila de totales, un grafico sobre los datos, formato condicional,
validacion de datos, un nombre definido y otras hojas.
"""
import datetime
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableFormula, TableStyleInfo

AQUI = os.path.dirname(os.path.abspath(__file__))


def plantilla_tabla():
    """Los datos van en una tabla de Excel, que crece; el resumen y el grafico la siguen."""
    wb = Workbook()
    portada = wb.active
    portada.title = 'Portada'
    portada['A1'] = 'Reporte mensual de ventas'
    portada['A1'].font = Font(bold=True, size=16)
    portada['A3'] = 'Preparado para: Tiendas del Norte'

    datos = wb.create_sheet('Datos')
    datos.append(['Tienda', 'Importe', 'Fecha', 'IVA'])
    muestra = [('ejemplo-1', 100.0, datetime.date(2026, 1, 1)), ('ejemplo-2', 200.0, datetime.date(2026, 1, 2))]
    for i, (t, imp, f) in enumerate(muestra, start=2):
        datos.append([t, imp, f, f'=Ventas[[#This Row],[Importe]]*0.16'])
        datos.cell(row=i, column=2).number_format = '"$"#,##0.00'
        datos.cell(row=i, column=3).number_format = 'dd/mm/yyyy'
        datos.cell(row=i, column=4).number_format = '"$"#,##0.00'
    # Fila de totales de la tabla
    datos['A4'] = 'Total'
    datos['B4'] = '=SUBTOTAL(109,Ventas[Importe])'
    datos['D4'] = '=SUBTOTAL(109,Ventas[IVA])'
    tabla = Table(displayName='Ventas', ref='A1:D4', totalsRowCount=1)
    tabla.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    tabla._initialise_columns()
    for col, nombre in zip(tabla.tableColumns, ['Tienda', 'Importe', 'Fecha', 'IVA']):
        col.name = nombre
    tabla.tableColumns[0].totalsRowLabel = 'Total'
    tabla.tableColumns[1].totalsRowFunction = 'sum'
    tabla.tableColumns[3].totalsRowFunction = 'sum'
    tabla.tableColumns[3].calculatedColumnFormula = TableFormula(attr_text='Ventas[[#This Row],[Importe]]*0.16')
    datos.add_table(tabla)
    datos.column_dimensions['A'].width = 18
    datos.column_dimensions['B'].width = 14
    datos.column_dimensions['C'].width = 12
    # Formato condicional y validacion sobre las filas de datos de muestra
    datos.conditional_formatting.add('B2:B3', CellIsRule(operator='greaterThan', formula=['1000'], fill=PatternFill(fill_type='solid', fgColor='FFC6EFCE')))
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0')
    dv.add('B2:B3')
    datos.add_data_validation(dv)

    resumen = wb.create_sheet('Resumen')
    resumen['A1'] = 'Total vendido'
    resumen['B1'] = '=SUM(Ventas[Importe])'
    resumen['A2'] = 'Tiendas'
    resumen['B2'] = '=ROWS(Ventas[Tienda])'
    resumen['A3'] = 'Importes (nombre definido)'
    resumen['B3'] = '=SUM(Importes)'
    grafico = BarChart()
    grafico.title = 'Ventas por tienda'
    grafico.add_data(Reference(datos, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    grafico.set_categories(Reference(datos, min_col=1, min_row=2, max_row=3))
    resumen.add_chart(grafico, 'D2')

    wb.defined_names['Importes'] = DefinedName('Importes', attr_text='Datos!$B$2:$B$3')
    notas = wb.create_sheet('Notas')
    notas['A1'] = 'Esta hoja no la toca AmoxSQL.'
    wb.save(os.path.join(AQUI, 'plantilla_tabla.xlsx'))


def plantilla_celda():
    """Sin tabla: los datos empiezan en B5, con encabezados ya escritos y una formula al lado."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Reporte'
    ws['B2'] = 'Ventas de la semana'
    ws['B2'].font = Font(bold=True, size=14)
    ws.append([])
    ws['B4'], ws['C4'], ws['D4'], ws['E4'] = 'Tienda', 'Unidades', 'Importe', 'Con IVA'
    for c in 'BCDE':
        ws[f'{c}4'].font = Font(bold=True)
    # Datos viejos de la semana anterior: tres filas que hay que limpiar si llegan menos.
    for i, (t, u, imp) in enumerate([('vieja-1', 1, 10.0), ('vieja-2', 2, 20.0), ('vieja-3', 3, 30.0)], start=5):
        ws[f'B{i}'], ws[f'C{i}'], ws[f'D{i}'] = t, u, imp
        ws[f'D{i}'].number_format = '#,##0.00'
        ws[f'E{i}'] = f'=D{i}*1.16'
    ws['G4'] = 'Nota fija a la derecha'
    ws['B20'] = 'Pie del reporte: no se toca'
    wb.save(os.path.join(AQUI, 'plantilla_celda.xlsx'))


if __name__ == '__main__':
    plantilla_tabla()
    plantilla_celda()
    print('ok')
